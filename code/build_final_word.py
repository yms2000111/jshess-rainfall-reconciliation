"""Build complete reading and review Word manuscripts from the final analysis."""
from common import ROOT,REGIONS
import ast,re,json,hashlib,shutil
import numpy as np,pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from final_equations import EQUATIONS

OUT=ROOT/'revision/submission';OUT.mkdir(exist_ok=True)
E=ROOT/'revision/extension';F=ROOT/'revision/final_figures'
text=(ROOT/'revision/manuscript_final.md').read_text(encoding='utf-8')
assert not re.search(r'TO_INSERT|TODO|PLACEHOLDER|\bTBC\b|\bTBD\b',text)
EQ_ORDER=[int(n) for n in re.findall(r'\[\[EQ(\d+)\]\]',text)]
DISPLAY={n:i+1 for i,n in enumerate(EQ_ORDER)}

# Reuse the already inspected native-OMML and typography functions, without
# executing the historical manuscript builder or reading its result tables.
old=ast.parse((ROOT/'code/build_publication.py').read_text(encoding='utf-8'))
functions=[n for n in old.body if isinstance(n,ast.FunctionDef) and n.name in ['word_math','document','table']]
exec(compile(ast.Module(body=functions,type_ignores=[]),'publication_helpers','exec'))
transform=etree.XSLT(etree.parse(r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL'))
original_math=word_math
original_table=table
def table(*args,**kwargs):
    original_table(*args,**kwargs)
    # The table already supplies cell padding; an extra empty paragraph can
    # spill onto a blank page immediately before the next page-break caption.
    last=args[0].paragraphs[-1]
    if not last.text:last._p.getparent().remove(last._p)
def word_math(n):
    result=original_math(n)
    # Bind each summand to Word's n-ary argument to remove blank math fields.
    for parent in reversed(list(result.iter())):
        for node in reversed(list(parent)):
            if node.tag!=qn('m:nary'):continue
            arg=node.find(qn('m:e'));nxt=node.getnext();anchor=node
            while nxt is None and anchor.getparent() is not None and anchor.getparent().tag!=qn('m:oMath'):
                anchor=anchor.getparent();nxt=anchor.getnext()
            if arg is not None and len(arg)==0 and nxt is not None:arg.append(nxt)
            prop=node.find(qn('m:naryPr'))
            limit=prop.find(qn('m:limLoc'))
            if limit is not None:limit.set(qn('m:val'),'undOvr')
    assert all(len(node.find(qn('m:e'))) for node in result.iter(qn('m:nary')))
    list(result.iter(qn('m:t')))[-1].text='\u2003('+str(DISPLAY[n])+')'
    return result

def paragraphs(doc,content):
    references=False
    for block in content.strip().split('\n\n'):
        block=block.strip()
        if block.startswith('# '):doc.add_paragraph(block[2:],style='Title')
        elif block.startswith('### '):doc.add_paragraph(block[4:],style='Heading 2')
        elif block.startswith('## '):
            p=doc.add_paragraph(block[3:],style='Heading 1')
            references=block=='## References'
            if block=='## 1 Introduction':p.paragraph_format.page_break_before=True
        elif block.startswith('[[EQ'):
            n=int(re.search(r'EQ(\d+)',block).group(1));p=doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(12)
            p.paragraph_format.keep_together=True
            p._p.append(word_math(n))
        else:
            p=doc.add_paragraph(block)
            if block in ['Musong Yang and Ce Yang','Beijing Institute of Technology, Beijing, China'] or block.startswith('Corresponding author:'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)
            if block=='Musong Yang and Ce Yang':p.runs[0].bold=True
            if block.startswith('Keywords:'):p.runs[0].italic=True
            if references:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
                p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(8)
                p.paragraph_format.keep_together=True
                for run in p.runs:run.font.size=Pt(11.5)

def figure(doc,name,caption):
    p=doc.add_paragraph();p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(10)
    p.add_run().add_picture(str(F/(name+'.png')),width=Inches(6.8))
    doc.add_paragraph(caption,style='Caption')

S=pd.read_csv(E/'scores.csv');B=pd.read_csv(E/'primary_comparisons.csv')
def score(k,m,col,lead='D28'):
    return S.query('domain==@k and method==@m and lead==@lead').iloc[0][col]

CAPTIONS=[
('Fig01_domains','Fig. 1. (a) Four geographical verification domains and (b) observed mean monthly rainfall during 2000 to 2019. The outlined boxes indicate the domain boundaries. The legend gives mean annual rainfall. The North arrow is separate from the northern-domain code N. AGCD values are area-weighted land averages; the Natural Earth outline is for orientation. The scale bar is evaluated at 41°S.'),
('Fig02_design','Fig. 2. (a) Weekly, fortnightly and 28-day aggregation structure. (b) Factorial experiment combining rainfall-only (H0) or climate-informed aggregate forecasts (H1) with five or seven hierarchy nodes. Weekly distributions are identical before reconciliation. (c) One outer validation fold, illustrated for 2010. Orange denotes the held-out year, grey its excluded neighbours and blue permitted fitting years. Inner year blocks estimate covariance; recalibration uses a further nested level.'),
('Fig03_weekly_skill','Fig. 3. Weekly CRPSS relative to the training seasonal climatology, pooled over 1060 held-out origins per domain during 2000 to 2019. Positive values indicate skill. Bottom-up forecasts retain the calibrated weekly marginals. H0 and H1 use seven-node nonnegative reconciliation. Open symbols identify forecast methods; the dotted line marks zero skill.'),
('Fig04_factorial','Fig. 4. Changes in 28-day CRPS for three factorial comparisons within each domain. H1 versus H0 holds seven nodes fixed; seven versus five nodes holds H1 fixed; H1 versus BU compares seven-node H1 with bottom-up. Changes are percentages of each domain’s out-of-fold climatological CRPS. Open symbols show mean differences and bars give descriptive 95% intervals from 10000 common year-block bootstrap resamples. Negative values favour H1 or the larger hierarchy. The three domain-averaged primary tests are reported in Table 3.'),
('Fig05_seasons','Fig. 5. Seasonal 28-day CRPS changes from H0 to H1 within the seven-node hierarchy. Seasons refer to forecast origins: DJF, December to February; MAM, March to May; JJA, June to August; SON, September to November. Open points are annual-block mean differences and bars are descriptive 95% year-bootstrap intervals. Negative values favour climate-informed aggregates. Vertical scales differ among panels.'),
('Fig06_coverage','Fig. 6. Observed versus nominal coverage of central 28-day forecast intervals at seven probability levels. Each curve pools 1060 held-out origins per domain. The dotted diagonal represents nominal coverage; points above it indicate overcoverage. Curves compare ECMWF (QC), bottom-up and seven-node H0 and H1 reconciliation. Interval widths at the 80% level are given in Table 4.'),
('Fig07_mechanism','Fig. 7. Covariance-predicted versus observed seasonal reductions in 28-day mean-square error relative to bottom-up. (a) Linear reconciliation and (b) means of nonnegative reconciled samples. Each open symbol represents a domain-season combination; colours identify domains and shapes identify seasons. The diagonal indicates agreement. Spearman correlations are 0.41 (95% year-bootstrap interval −0.03 to 0.61) and 0.34 (−0.06 to 0.47), respectively.'),
('Fig08_recalibration','Fig. 8. CRPS changes after coherent power-map recalibration of seven-node H1 trajectories. Positive changes indicate deterioration. Open points show annual-block mean changes and bars give descriptive 95% year-bootstrap intervals. W1 to W4 are weekly accumulations and 28 d is the total. Annotations show the change in 28-day coverage of central 80% intervals. The same nonnegative proportional adjustment is applied to every week within each trajectory.')]

rows1=[[k,r['name'],f'{r["lon"][0]:.2f} to {r["lon"][1]:.2f}°E',f'{abs(r["lat"][1]):.2f} to {abs(r["lat"][0]):.2f}°S'] for k,r in REGIONS.items()]
methods=['Raw','Bottom_up','H0_Direct','H1_Direct','H0_5_NN','H0_7_NN','H1_5_NN','H1_7_NN']
labels=['ECMWF (QC)','Bottom-up','H0 direct','H1 direct','H0, five nodes','H0, seven nodes','H1, five nodes','H1, seven nodes']
rows2=[[label]+[f'{score(k,m,"crps"):.3f}' for k in REGIONS] for m,label in zip(methods,labels)]
baseline_labels={'H0_7_NN':'H1 vs H0, seven nodes','H1_5_NN':'Seven vs five nodes, H1','Bottom_up':'H1, seven nodes vs bottom-up'}
rows3=[[baseline_labels[row.baseline],f'{row.delta_normalized_percent:+.3f}',f'{row.lower:+.3f} to {row.upper:+.3f}',f'{row.p_holm:.3f}',f'{row.better_years}/20'] for row in B.itertuples()]
rows4=[[k,lab,f'{score(k,m,"coverage80")*100:.1f}',f'{score(k,m,"width80"):.2f}',f'{score(k,m,"twcrps"):.3f}',f'{score(k,m,"brier90"):.3f}',str(int(score(k,m,'events90')))] for k in REGIONS for m,lab in [('Bottom_up','Bottom-up'),('H0_7_NN','H0, seven'),('H1_7_NN','H1, seven')]]
TABLES=[
('Table 1. Verification-domain boundaries. Each domain contains four archived forecast cells. Observations use valid AGCD land cells within the same bounds.',['Code','Domain','Longitude','Latitude'],rows1,[.5,2.5,1.9,1.9]),
('Table 2. Held-out 28-day CRPS (mm), evaluated over 1060 origins per domain during 2000 to 2019. Lower scores indicate better forecasts. H0 uses ensemble rainfall predictors; H1 adds lagged climate indices only at aggregate scales. All reconciled forecasts use the nonnegative sample projection.',['Method','N','NE','SE','SW'],rows2,[2.4,1.1,1.1,1.1,1.1]),
('Table 3. Three primary factorial comparisons at 28 days. Effects and intervals are percentages of climatological CRPS, with equal domain weights. Negative values favour the first method. Intervals use year-block bootstrapping; p values use annual sign flips with Holm adjustment. Better years counts negative annual differences.',['Comparison','Effect (%)','95% interval (%)','Adjusted p','Better years'],rows3,[2.55,.85,1.65,.85,.9]),
('Table 4. Held-out 28-day uncertainty and upper-tail diagnostics. Coverage and width refer to central 80% intervals. Threshold-weighted CRPS and Brier scores use the seasonal training 90th percentile; event counts are observed exceedances among 1060 origins. H0 and H1 have seven nodes.',['Domain','Method','Coverage (%)','Width (mm)','Tail CRPS (mm)','Brier score','Events'],rows4,[.55,1.45,.9,.9,1.0,.9,1.1])]

doc=document();doc.core_properties.title=text.splitlines()[0][2:];doc.core_properties.author='Musong Yang; Ce Yang'
paragraphs(doc,text)
for i,args in enumerate(TABLES):table(doc,*args,newpage=i in [0,2])
for name,caption in CAPTIONS:figure(doc,name,caption)
doc.save(OUT/'JSHESS_manuscript_reading.docx')
for sec in doc.sections:
    ln=OxmlElement('w:lnNumType')
    for key,value in [('countBy','1'),('start','1'),('restart','continuous'),('distance','160')]:ln.set(qn('w:'+key),value)
    sec._sectPr.append(ln)
doc.save(OUT/'JSHESS_manuscript_submission.docx')

# Supplementary tables referenced by the complete main text.
supp=document();supp.core_properties.title='Supplementary material';supp.core_properties.author='Musong Yang; Ce Yang'
paragraphs(supp,'# Supplementary material\n\n'+text.splitlines()[0][2:]+'\n\nMusong Yang and Ce Yang\n\nAll scores use the same 20 outer validation years as the main manuscript. Bootstrap intervals are descriptive. Differences are first method minus the stated reference, so negative CRPS differences are favourable.')
q=pd.read_csv(ROOT/'revision/qc_by_week.csv')
qrows=[[r.domain,f'W{r.week}',str(r.negative),f'{r.percent:.2f}',f'{r.minimum_mm:.3f}'] for r in q.itertuples()]
table(supp,'Table S1a. Negative weekly archive increments before correction. Each domain-week contains 11660 member-origin increments.',['Domain','Week','Negative count','Percentage','Minimum (mm)'],qrows,[.8,.7,1.6,1.5,2.2])
qrows=[]
for qc,label in [('primary','Isotonic'),('positive','Positive-part'),('signed','Signed predictors')]:
    c=pd.read_csv(ROOT/f'revision/{qc}/primary_comparisons.csv')
    # Retain all three prespecified contrasts with their original column names.
    c.to_csv(OUT/f'quality_control_{qc}_primary_comparisons.csv',index=False)
    z=pd.read_csv(ROOT/f'revision/{qc}/comparisons.csv').query("method=='MinT_NN' and baseline=='Bottom_up' and lead=='D28'")
    for row in z.itertuples():qrows.append([label,row.domain,f'{row.delta:+.3f}',f'{row.lower:+.3f} to {row.upper:+.3f}'])
table(supp,'Table S1b. Five-node homogeneous reconciliation minus bottom-up 28-day CRPS under the three quality-control treatments.',['Treatment','Domain','Difference (mm)','95% interval (mm)'],qrows,[2.1,.7,1.5,2.5],newpage=True)
P=pd.read_csv(ROOT/'revision/primary/scores.csv');PC=pd.read_csv(ROOT/'revision/primary/comparisons.csv')
rows=[]
for k in REGIONS:
    a=P.query("domain==@k and lead=='D28'").set_index('method')
    d=PC.query("domain==@k and lead=='D28' and method=='ECC11_NN' and baseline=='Schaake11_NN'").iloc[0]
    rows.append([k,f'{a.loc["ECC11_NN","crps"]:.3f}',f'{a.loc["Schaake11_NN","crps"]:.3f}',f'{d.delta:+.3f}',f'{d.lower:+.3f} to {d.upper:+.3f}'])
table(supp,'Table S2. Matched 11-sample dependence comparison for the homogeneous five-node forecast. Marginal quantiles are identical before reordering. Scores and paired differences are in millimetres.',['Domain','ECC CRPS','Schaake CRPS','ECC minus Schaake','95% interval'],rows,[.65,1.1,1.2,1.55,2.3],newpage=True)
rows=[[k,l,lab,f'{score(k,m,"crps",l):.3f}',f'{score(k,m,"coverage80",l)*100:.1f}'] for k in REGIONS for l in ['F12','F34'] for m,lab in [('Bottom_up','Bottom-up'),('H0_7_NN','H0, seven'),('H1_7_NN','H1, seven')]]
table(supp,'Table S3. Fortnightly CRPS and central 80% interval coverage. F12 spans Weeks 1 and 2; F34 spans Weeks 3 and 4.',['Domain','Period','Method','CRPS (mm)','Coverage (%)'],rows,[.65,.75,1.8,1.8,1.8],newpage=True)
J=pd.read_csv(E/'joint_scores_annual.csv');rng=np.random.default_rng(20261024);draw=rng.integers(0,20,(10000,20));rows=[]
for k in REGIONS:
    a=J.query('domain==@k').groupby('method')[['energy','variogram']].mean()
    for m,lab in [('Raw','ECMWF (QC)'),('Bottom_up','Bottom-up'),('H0_7_NN','H0, seven'),('H1_7_NN','H1, seven')]:rows.append([k,lab,f'{a.loc[m,"energy"]:.3f}',f'{a.loc[m,"variogram"]:.3f}'])
table(supp,'Table S4. Joint four-week forecast scores. The energy score uses Euclidean distance; the variogram score uses order 0.5 and equal weights for the six unique week pairs. Lower values are better. The two scores have different scales and are compared only within columns.',['Domain','Method','Energy score','Variogram score'],rows,[.7,2.3,1.9,1.9],newpage=True)
st=pd.read_csv(E/'sample_stability.csv').query("lead=='D28'").groupby(['domain','size','seed']).delta_bu.mean().reset_index()
rows=[]
for (k,size),a in st.groupby(['domain','size'],sort=False):rows.append([k,str(size),f'{a.delta_bu.mean():+.3f}',f'{a.delta_bu.min():+.3f} to {a.delta_bu.max():+.3f}'])
table(supp,'Table S5. Sampling sensitivity of seven-node H1 minus bottom-up 28-day CRPS. For each size, the full 20-year mean is first calculated for each of three seeds. The seed range describes numerical sampling variation, not a confidence interval.',['Domain','Samples','Mean difference (mm)','Range across seeds (mm)'],rows,[.7,.9,2.2,3.0],newpage=True)
RS=pd.read_csv(E/'recalibration/scores.csv');RC=pd.read_csv(E/'recalibration/comparisons.csv');rows=[]
for k in REGIONS:
    for lead in ['W1','W2','W3','W4','D28']:
        a=RS.query('domain==@k and lead==@lead').set_index('method');d=RC.query("domain==@k and lead==@lead and measure=='crps'").iloc[0]
        rows.append([k,lead,f'{d.delta:+.3f}',f'{d.lower:+.3f} to {d.upper:+.3f}',f'{a.loc["Original","coverage80"]*100:.1f}',f'{a.loc["Coherent_recalibration","coverage80"]*100:.1f}'])
table(supp,'Table S6. Coherent recalibration minus original seven-node H1. CRPS differences are in millimetres. Coverage columns show central 80% intervals before and after recalibration.',['Domain','Scale','CRPS change','95% interval','Before (%)','After (%)'],rows,[.65,.65,1.15,2.05,1.15,1.15],newpage=True)
figure(supp,'FigS01_tail_reliability','Fig. S1. Reliability for 28-day probabilities exceeding the seasonal training 90th percentile. Open points represent ten probability bins, omitting bins with fewer than 20 cases. Vertical bars are descriptive 95% year-bootstrap intervals. The diagonal denotes agreement between predicted probabilities and observed frequencies.')
figure(supp,'FigS02_pit','Fig. S2. Randomised probability integral transform histograms for 28-day distributions. Each panel pools 1060 origins; the dotted horizontal line is the uniform-bin reference of 0.1. The step curves show the observed proportion in each bin.')
supp.save(OUT/'JSHESS_supplementary_material.docx')

statement='Rainfall forecasts for individual weeks and longer periods need to give farmers and water managers a consistent picture. In four Australian areas, adding climate information improved some northern forecasts, while simply adding more accumulation periods brought little benefit. These results show why forecast services should compare a combined forecast with the sum of its weekly forecasts and check the uncertainty of both.'
(OUT/'statement_of_importance.txt').write_text(statement+'\n',encoding='utf-8')
(OUT/'figure_captions.txt').write_text('\n\n'.join(c for _,c in CAPTIONS),encoding='utf-8')
abstract=text.split('## Abstract\n\n')[1].split('\n\nKeywords:')[0]
main=text.split('## 1 Introduction')[1].split('## References')[0]
audit={'abstract_words':len(abstract.split()),'main_and_captions_words':len(main.split())+sum(len(c.split()) for _,c in CAPTIONS)+sum(len(x[0].split()) for x in TABLES),'references':len(text.split('## References')[1].strip().split('\n\n')),'equations':len(EQ_ORDER),'main_figures':len(CAPTIONS),'main_tables':len(TABLES),'importance_words':len(statement.split()),'body_points':12,'headings_points':[20,16,14],'figure_width_inches':6.8,'reading_line_numbers':False,'submission_line_numbers':True,'equation_source_sha256':hashlib.sha256(json.dumps(EQUATIONS,sort_keys=True).encode()).hexdigest()}
assert audit['abstract_words']<=250 and audit['main_and_captions_words']<=7500 and 50<=audit['importance_words']<=80
(OUT/'format_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print(json.dumps(audit,indent=2))
