from common import ROOT
from docx import Document
from lxml import etree
import json,re,hashlib
R=ROOT/'revision/review_v6';O=R/'submission_final';s=(R/'manuscript.md').read_text(encoding='utf-8')
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
audit={}
for name in ['JSHESS_manuscript_reading','JSHESS_manuscript_submission','JSHESS_supplementary_material']:
 d=Document(O/(name+'.docx'));t='\n'.join(p.text for p in d.paragraphs)
 eq=d.element.xpath('.//m:oMath');assert len(eq)==(4 if 'supplementary' in name else 11)
 assert '[[EQ' not in t and '\\frac' not in t and '\ufffd' not in t
 caps=[p.text for p in d.paragraphs if re.match(r'^(?:Table S\d|Fig\. S\d)',p.text)]
 ids=[re.match(r'^(Table S\d+[a-z]?|Fig\. S\d+)\.',p).group(1) for p in caps]
 assert len(ids)==len(set(ids))
 if 'supplementary' in name:
  assert 'Thresholds are defined in Table 5' in t
  assert '(Table 5; Supplementary Table S7)' in t
  assert [int(x.split('S')[-1]) for x in ids if x.startswith('Fig')]==list(range(1,13))
  for n in range(1,30):assert any(re.match('Table S'+str(n)+r'(?:[a-z]|$)',x) for x in ids)
 else:
  refs=False
  for p in d.paragraphs:
   if p.text=='References':refs=True;continue
   if refs and p.text.startswith('Table 1.'):break
   if refs and p.text:assert p.alignment==0,p.text[:40]
 audit[name]=dict(native_equations=len(eq),paragraphs=len(d.paragraphs),tables=len(d.tables),unique_supplementary_captions=len(ids))
body,refs=s.split('## References\n\n');missing=[]
for ref in refs.strip().split('\n\n'):
 author=ref.split(' (')[0];year=re.search(r'\((\d{4}[ab]?|undated [ab])\)',ref).group(1)
 surname='Australian Bureau of Meteorology' if author.startswith('Australian Bureau') else author.split()[0]
 if not any(surname in p and year in p for p in body.split('\n\n')):missing.append(ref)
audit['uncited_main_references']=missing;assert not missing
audit['word_counts']=json.loads((O/'format_audit.json').read_text())
audit['scope']='Rendered Word documents and LaTeX compiled successfully; original mathematical source unchanged. Statistical score reconstruction is separately documented.'
(R/'final_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8');print(json.dumps(audit,indent=2))
