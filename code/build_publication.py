"""Create matching Word and LaTeX manuscripts from one text and equation source."""
from common import *
import re,json,hashlib,shutil,pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from equations import EQUATIONS
OUT=ROOT/'submission';OUT.mkdir(exist_ok=True);TEX=ROOT/'latex';(TEX/'images').mkdir(exist_ok=True)
text=(ROOT/'manuscript.md').read_text(encoding='utf-8');head,back=text.split('## Data availability',1);back='## Data availability'+back
S=pd.read_csv(ROOT/'results/scores.csv')
CAPTIONS=[
('Fig01_regions','Fig. 1. (a) Four regional verification domains and (b) AGCD monthly rainfall climatologies for 2000–2012. Boxes show the exact verification support. N, NE, SE and SW denote northern, northeastern, southeastern and southwestern domains. The geographical outline is from Natural Earth; the scale bar is evaluated at 41°S.'),
('Fig02_design','Fig. 2. (a) Chronological fitting, calibration and independent-test partitions. Origin counts exclude verification windows crossing fitting or calibration boundaries. XGBoost tuning is nested within the fitting period. (b) Construction of the EMOS forecast trajectories and the four-week aggregation identity.'),
('Fig03_skill','Fig. 3. Weekly CRPSS relative to fitting-period seasonal climatology, evaluated over 265 origins in 2015–2019 in each region. Positive values indicate skill. Direct and bottom-up EMOS have identical weekly forecasts, so their curves overlap. Open symbols and line styles distinguish methods; ECMWF (QC) includes the documented archive quality control.'),
('Fig04_reconciliation','Fig. 4. Paired CRPS changes for MinT-NN relative to direct EMOS and bottom-up EMOS. Negative values indicate improvement. Open symbols show sample means; bars show descriptive 95% percentile intervals from 5000 resamples of the five complete test-origin years. Weekly changes relative to the two baselines are identical. All panels use the same horizontal scale.'),
('Fig05_reliability','Fig. 5. Reliability for probabilities of exceeding the fitting-period seasonal upper tercile, pooling Weeks 2–4. The dotted diagonal indicates perfect reliability. Ten probability bins are used; bins with fewer than 20 cases are omitted. Each regional pool contains 795 overlapping origin–lead pairs. Direct and bottom-up EMOS coincide at weekly scales.'),
('Fig06_coverage','Fig. 6. Coverage of central 80% prediction intervals at each accumulation scale. The dotted line marks nominal coverage. The 28-day total is a separate aggregate product rather than an additional weekly lead. Each point uses 265 test origins; direct and bottom-up EMOS coincide for Weeks 1–4.')]

def score(k,m,col):return S.query("region==@k and method==@m and lead=='D28'").iloc[0][col]
rows1=[[k,r['name'],f'{r["lon"][0]:.2f}–{r["lon"][1]:.2f}°E',f'{abs(r["lat"][1]):.2f}–{abs(r["lat"][0]):.2f}°S'] for k,r in REGIONS.items()]
methods=['Climatology','Raw','Seasonal BC','Quantile mapping','CSG EMOS','EMOS bottom-up','EMOS MinT-NN','OLS','XGBoost']
labels=['Climatology','ECMWF (QC)','Seasonal correction','Quantile mapping','Direct EMOS','EMOS bottom-up','EMOS MinT-NN','OLS','XGBoost']
rows2=[[l]+[f'{score(k,m,"crps"):.2f}' for k in REGIONS] for m,l in zip(methods,labels)]
rows3=[[k,l,f'{100*score(k,m,"coverage80"):.1f}',f'{score(k,m,"width80"):.2f}',f'{score(k,m,"coherence_mae"):.2f}'] for k in REGIONS for m,l in [('CSG EMOS','Direct EMOS'),('EMOS bottom-up','Bottom-up'),('EMOS MinT-NN','MinT-NN')]]
TABLES=[('Table 1. Regional verification domains. Each contains four archived forecast cells; observations are averaged over valid AGCD land cells within the same bounds.',['Code','Location','Longitude','Latitude'],rows1,[.55,2.65,1.8,1.8]),('Table 2. Independent-test 28-day CRPS (mm). Each estimate uses 265 origins in 2015–2019. Lower scores indicate better forecasts.',['Method','N','NE','SE','SW'],rows2,[2.4,1.1,1.1,1.1,1.1]),('Table 3. Independent-test 28-day interval coverage, mean interval width and mean absolute sample-wise aggregation discrepancy. Intervals have nominal coverage of 80%.',['Region','Method','Coverage (%)','Width (mm)','Discrepancy (mm)'],rows3,[.65,1.7,1.4,1.4,1.65])]
transform=etree.XSLT(etree.parse(r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL'))

def word_math(n):
 ns={'mml':'http://www.w3.org/1998/Math/MathML'}
 ml=etree.fromstring(convert(EQUATIONS[n],display='block').encode())
 for e in ml.xpath('//mml:mover',namespaces=ns):
  if ''.join(e[-1].itertext()) in ['^','~','¯','‾']:
   e.set('accent','true')
 for e in ml.xpath('//mml:mspace',namespaces=ns):
  t=etree.Element('{'+ns['mml']+'}mtext');t.text='\u2003' if e.get('width')!='2em' else '\u2003\u2003';e.getparent().replace(e,t)
 root=transform(ml).getroot()
 for matrix in list(root.iter(qn('m:m'))):
  parent=matrix.getparent();idx=parent.index(matrix)
  if idx>0:
   for tx in parent[idx-1].iter(qn('m:t')):
    if tx.text and tx.text.endswith('('):tx.text=tx.text[:-1]
  if idx+1<len(parent):
   for tx in parent[idx+1].iter(qn('m:t')):
    if tx.text and tx.text.startswith(')'):tx.text=tx.text[1:]
  delimiter=OxmlElement('m:d');pr=OxmlElement('m:dPr')
  for key,val in [('begChr','('),('endChr',')'),('grow','1')]:
   el=OxmlElement('m:'+key);el.set(qn('m:val'),val);pr.append(el)
  delimiter.append(pr);arg=OxmlElement('m:e');parent.replace(matrix,delimiter);arg.append(matrix);delimiter.append(arg)
 for grow in root.iter(qn('m:grow')):
  if grow.getparent().tag==qn('m:naryPr'):grow.set(qn('m:val'),'0')
 # Word's default body font is not a mathematical font. Set the math runs explicitly.
 for r in root.iter(qn('m:r')):
  pr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
  for a in ['ascii','hAnsi','cs']:rf.set(qn('w:'+a),'Cambria Math')
  pr.append(rf);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');pr.append(sz);r.insert(0,pr)
  tx=''.join(r.itertext())
  if any(v in tx for v in ['CRPS','Gamma','arg','min','max']):
   mp=r.find(qn('m:rPr'))
   if mp is None:mp=OxmlElement('m:rPr');r.insert(0,mp)
   sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');mp.append(sty)
 # Attach adjacent nested sums to their argument, avoiding an empty n-ary placeholder.
 for parent in list(root.iter()):
  children=list(parent)
  for i,e in enumerate(children[:-1]):
   if e.tag==qn('m:nary') and children[i+1].tag==qn('m:nary'):
    arg=e.find(qn('m:e'))
    if arg is not None and len(arg)==0:arg.append(children[i+1])
 number=OxmlElement('m:r');pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);number.append(pr);tx=OxmlElement('m:t');tx.text='\u2003('+str(n)+')';number.append(tx);root.append(number)
 para=OxmlElement('m:oMathPara');para.append(root);return para

def document():
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.85)
 for n in ['Normal','Title','Heading 1','Heading 2','Caption']:
  st=doc.styles[n];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);rf=st.element.get_or_add_rPr().get_or_add_rFonts()
  for attr in list(rf.attrib):rf.attrib.pop(attr)
  for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
 for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
 normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.3;normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;normal.paragraph_format.widow_control=True
 st=doc.styles['Title'];st.font.size=Pt(20);st.font.bold=True;st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;st.paragraph_format.space_after=Pt(16);st.paragraph_format.line_spacing=1.1
 for n,size in [('Heading 1',16),('Heading 2',14)]:
  st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(9);st.paragraph_format.keep_with_next=True
 st=doc.styles['Caption'];st.font.size=Pt(11);st.font.italic=False;st.font.bold=False;st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(9)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 return doc

def paragraphs(doc,content):
 reference=False
 for block in content.strip().split('\n\n'):
  block=block.strip()
  if block.startswith('# '):doc.add_paragraph(block[2:],style='Title')
  elif block.startswith('### '):doc.add_paragraph(block[4:],style='Heading 2')
  elif block.startswith('## '):
   doc.add_paragraph(block[3:],style='Heading 1');reference=block=='## References'
  elif block.startswith('[[EQ'):
   n=int(re.search(r'EQ(\d+)',block).group(1));p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(12);p.paragraph_format.line_spacing=1.3
   p._p.append(word_math(n))
  else:
   p=doc.add_paragraph(block)
   if block.startswith('[') or block.startswith('Corresponding author:'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)
   if block.startswith('Keywords:'):p.runs[0].italic=True
   if reference:p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(9)

def table(doc,caption,headers,rows,widths,newpage=False):
 p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True;p.paragraph_format.page_break_before=newpage
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for cell,txt,w in zip(t.rows[0].cells,headers,widths):cell.text=txt;cell.width=Inches(w)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  cells=t.add_row().cells
  for cell,txt,w in zip(cells,row,widths):cell.text=str(txt);cell.width=Inches(w)
 for ir,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ic,cell in enumerate(row.cells):
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   pr.append(bd)
   for p in cell.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ic<2 else WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
    for run in p.runs:run.font.size=Pt(11);run.bold=ir==0
 doc.add_paragraph().paragraph_format.space_after=Pt(3)

doc=document();paragraphs(doc,head)
for i,args in enumerate(TABLES):table(doc,*args,newpage=i in [0,2])
for name,caption in CAPTIONS:
 p=doc.add_paragraph();p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(10);p.add_run().add_picture(str(ROOT/f'figures/{name}.png'),width=Inches(6.8));doc.add_paragraph(caption,style='Caption')
start=len(doc.paragraphs);paragraphs(doc,back);doc.paragraphs[start].paragraph_format.page_break_before=True
doc.core_properties.title=head.splitlines()[0][2:];doc.core_properties.author='';doc.save(OUT/'JSHESS_reading.docx')
for sec in doc.sections:
 ln=OxmlElement('w:lnNumType');ln.set(qn('w:countBy'),'1');ln.set(qn('w:start'),'1');ln.set(qn('w:restart'),'continuous');ln.set(qn('w:distance'),'160');sec._sectPr.append(ln)
doc.save(OUT/'JSHESS_submission.docx')

mapping={'−':r'\ensuremath{-}','±':r'\ensuremath{\pm}','×':r'\ensuremath{\times}','²':r'\textsuperscript{2}','⁻':r'\textsuperscript{-}','⁸':r'\textsuperscript{8}','ᵀ':r'\ensuremath{^{\mathsf{T}}}','°':r'\ensuremath{^\circ}','μ':r'\ensuremath{\mu}','σ':r'\ensuremath{\sigma}','δ':r'\ensuremath{\delta}','κ':r'\ensuremath{\kappa}','θ':r'\ensuremath{\theta}','λ':r'\ensuremath{\lambda}','ŷ':r'\ensuremath{\widehat y}','–':'--','—':'---','≤':r'\ensuremath{\leq}','≥':r'\ensuremath{\geq}','→':r'\ensuremath{\to}','‘':'`','’':"'"}
def esc(s):
 out=''
 for ch in s:
  out+=mapping.get(ch,('\\'+ch if ch in '&%$#_{}' else ch))
 return out
def prose(s):
 parts=re.split(r'(https?://\S+)',s);return ''.join(r'\url{'+v+'}' if v.startswith('http') else esc(v) for v in parts)
def texparas(content):
 out=[]
 for block in content.strip().split('\n\n'):
  if block.startswith('# '):out.append(r'\begin{center}{\fontsize{20}{24}\selectfont\bfseries '+esc(block[2:])+r'}\end{center}\vspace{5pt}')
  elif block.startswith('### '):out.append(r'\subsection*{'+esc(block[4:])+'}')
  elif block.startswith('## '):out.append(r'\section*{'+esc(block[3:])+'}')
  elif block.startswith('[[EQ'):
   n=int(re.search(r'EQ(\d+)',block).group(1));out.append(r'\begin{equation}'+EQUATIONS[n]+r'\tag{'+str(n)+r'}\end{equation}')
  elif block.startswith('[') or block.startswith('Corresponding author:'):out.append(r'\begin{center}'+esc(block)+r'\end{center}')
  else:out.append(prose(block))
 return '\n\n'.join(out)
preamble=r'''\documentclass[12pt,letterpaper]{article}
\usepackage[margin=0.85in]{geometry}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{mathptmx,amsmath,amssymb,graphicx,longtable,array,booktabs,caption,titlesec,setspace,xurl,hyperref,lineno}
\hypersetup{hidelinks}
\setstretch{1.3}
\setlength{\parindent}{0pt}
\setlength{\parskip}{7pt}
\titleformat{\section}{\fontsize{16}{19}\selectfont\bfseries}{}{0pt}{}
\titleformat{\subsection}{\fontsize{14}{17}\selectfont\bfseries}{}{0pt}{}
\titlespacing*{\section}{0pt}{16pt}{9pt}
\titlespacing*{\subsection}{0pt}{16pt}{9pt}
\captionsetup{font=small,justification=raggedright,singlelinecheck=false}
\setlength{\emergencystretch}{3em}
\begin{document}
\ifdefined\reviewlinenos\linenumbers\fi
'''
tex=preamble+texparas(head)
for i,(caption,headers,rows,widths) in enumerate(TABLES):
 tex+='\n\\clearpage\n' if i in [0,2] else '\n'
 tex+='\\begin{small}\n\\begin{longtable}{'+''.join('p{'+f'{w-.17:.2f}'+'in}' for w in widths)+'}\n\\caption*{'+esc(caption)+'}\\\\\n\\toprule\n'+' & '.join(esc(h) for h in headers)+r'\\\midrule\endhead'+'\n'
 tex+='\n'.join(' & '.join(esc(str(v)) for v in row)+r'\\' for row in rows)+'\n\\bottomrule\n\\end{longtable}\n\\end{small}\n'
for name,caption in CAPTIONS:
 shutil.copy2(ROOT/f'figures/{name}.pdf',TEX/f'images/{name}.pdf');tex+='\n\\clearpage\n\\begin{center}\n\\includegraphics[width=\\linewidth]{images/'+name+'.pdf}\n\\captionof*{figure}{'+esc(caption)+'}\n\\end{center}\n'
tex+='\n\\clearpage\n'+texparas(back)+'\n\\end{document}\n';(TEX/'main.tex').write_text(tex,encoding='utf-8');(TEX/'review.tex').write_text('\\def\\reviewlinenos{1}\n\\input{main.tex}\n',encoding='utf-8')
(OUT/'figure_captions.txt').write_text('\n\n'.join(c for _,c in CAPTIONS),encoding='utf-8')
statement='Rainfall forecasts help Australian communities plan farming and water use several weeks ahead. We found that making weekly and four-week forecasts agree can improve some overall scores, but a simple sum of weekly forecasts sometimes works better than combining forecasts from both scales. The results show why forecast services should check the accuracy and uncertainty of each rainfall product alongside its consistency with other products.'
(OUT/'statement_of_importance.txt').write_text(statement+'\n',encoding='utf-8')
abstract=text.split('## Abstract\n\n')[1].split('\n\nKeywords:')[0]
audit=dict(abstract_words=len(abstract.split()),main_plus_captions_words=len(head.split('## 1 Introduction')[1].split())+sum(len(c.split()) for _,c in CAPTIONS)+sum(len(t[0].split()) for t in TABLES),importance_words=len(statement.split()),importance_sentences=3,references=len(back.split('## References')[1].strip().split('\n\n')),equations=len(EQUATIONS),equation_sha256=hashlib.sha256(json.dumps(EQUATIONS,sort_keys=True).encode()).hexdigest(),body_pt=12,title_pt=20,heading1_pt=16,heading2_pt=14,reading_line_numbers=False,submission_line_numbers=True)
assert audit['abstract_words']<=250 and audit['main_plus_captions_words']<=7500 and 50<=audit['importance_words']<=80
(OUT/'format_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8');print(audit)
