"""Deterministic cross-reference, equation, file and numerical release checks."""
from common import ROOT
from docx import Document
from docx.oxml.ns import qn
import re,json,hashlib
import numpy as np,pandas as pd
R=ROOT/'revision/review_v4';D=R/'submission';s=(R/'manuscript.md').read_text(encoding='utf-8')
a=Document(D/'JSHESS_manuscript_reading.docx');b=Document(D/'JSHESS_manuscript_submission.docx');su=Document(D/'JSHESS_supplementary_material.docx')
def content(doc):return '\n'.join([p.text for p in doc.paragraphs]+['|'.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
assert content(a)==content(b)
eq=list(a._element.iter(qn('m:oMath')));assert len(eq)==13
for e in eq:
 for n in e.iter(qn('m:nary')):assert len(n.find(qn('m:e')))>0
assert len(list(su._element.iter(qn('m:oMath'))))==2
assert not list(a._element.iter(qn('w:lnNumType')))
assert list(b._element.iter(qn('w:lnNumType')))
labels=[re.match(r'Table (S\d+[a-z]?)\.',p.text).group(1) for p in su.paragraphs if re.match(r'Table S\d+[a-z]?\.',p.text)]
assert len(labels)==len(set(labels))
assert [int(re.search(r'\d+',x).group()) for x in labels]==sorted(int(re.search(r'\d+',x).group()) for x in labels)
numbers=set(int(re.search(r'\d+',x).group()) for x in labels);assert numbers==set(range(1,22))
referred=set(map(int,re.findall(r'\bS(\d+)\b',s)))
for lo,hi in re.findall(r'S(\d+)–S(\d+)',s):referred.update(range(int(lo),int(hi)+1))
assert numbers<=referred,(numbers-referred)
assert not re.search(r'EXTENSION_|PLACEHOLDER|TO_INSERT|\bTODO\b|\bTBD\b|□|�',content(a)+content(su))
assert '—' not in s.split('## References')[0]
assert not re.search(r'Bottom_up|Weekly_climate_BU|H0_7_NN|Reviewer five years',content(su))
access_ref=pd.read_csv(R/'access_s2/comparisons.csv').query('lead=="D28"').set_index('baseline')
access_table=next(t for t in su.tables if [c.text for c in t.rows[0].cells]==['Season','Reference','Difference (mm)','95% interval'])
short={'Raw':'Raw','BU-rain':'Bottom_up','BU-clim':'Weekly_climate_BU','H0':'H0'}
for row in access_table.rows[1:]:
 cells=[c.text for c in row.cells]
 if cells[0]=='Annual':
  r=access_ref.loc[short[cells[1]]]
  assert cells[2]==f'{r.delta:+.3f}' and cells[3]==f'{r.lower:+.3f} to {r.upper:+.3f}'
for p in (R/'adaptive').glob('*.npz'):
 if p.name.startswith('inner_'):continue
 f=np.load(p);assert np.isfinite(f['scores']).all() and np.isfinite(f['probs']).all()
 for key in ['Constant','Seasonal','State','Equal']:
  w=f['weight_'+key];assert np.allclose(w.sum(1),1) and w.min()>=0
 meta=json.loads(p.with_suffix('.json').read_text());assert not set(meta['train'])&set(meta['test'])
 raw=np.load(ROOT/f'data/processed/{meta["domain"]}.npz');dates=pd.DatetimeIndex(raw['dates']);tr=np.array(meta['train']);te=np.array(meta['test'])
 if isinstance(meta['year'],int):assert not set(dates[tr].year)&{meta['year']-1,meta['year'],meta['year']+1}
 else:assert (dates[tr]+pd.Timedelta(days=28)).max()<dates[te].min()
 for mod in meta['models'].values():assert np.isfinite(mod['beta']).all()
mapping=[]
for label in labels:
 p=next(p for p in su.paragraphs if p.text.startswith('Table '+label+'.'));mapping.append(dict(label=label,caption=p.text))
(D/'supplementary_table_index.csv').write_text(pd.DataFrame(mapping).to_csv(index=False),encoding='utf-8')
audit=dict(main_equations=13,supplementary_equations=2,reading_submission_content_identical=True,supplementary_labels=labels,all_supplementary_numbers_referenced=True,mixture_files=len([p for p in (R/'adaptive').glob('*.npz') if not p.name.startswith('inner_')]),placebo_fits=len(list((R/'placebo').glob('*.npz'))),geographical_folds=len(list((R/'coverage').glob('*_20*.npz'))),access_folds=len(list((R/'access_s2/folds').glob('*.npz'))),word_limit=json.loads((D/'format_audit.json').read_text()),style_skill='https://github.com/Zink-Liu/Academic-Writing-Humanizer',style_review='Manual scientific and language checks using the installed skill; no automated AI detector score claimed.')
(D/'release_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8');print(json.dumps(audit,indent=2))
