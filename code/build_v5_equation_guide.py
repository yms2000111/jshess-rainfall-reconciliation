"""Chinese reading companion using the exact manuscript OMML equations."""
from common import ROOT
import runpy,json
ns=runpy.run_path(str(ROOT/'code/build_v5_word.py'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
items=[
(3,'1 降水概率分布',[
'用途：把连续的正降水分布与“无降水”的概率放在同一个模型中。',
'Y 是最终降水量；Z 是尚未经过移动和截零的伽马变量；κ 是形状参数，θ 是尺度参数，δ 是非负移动量。先计算 Z−δ，再与 0 取较大值。',
'当 Z≤δ 时，Y=0，因此零降水具有正概率。μ 和 σ² 是 Z 的均值和方差，不能直接当成 Y 的均值和方差。代码在训练降水尺度下拟合，再恢复为毫米。']),
(4,'2 集合特征如何决定分布',[
'用途：把动力模式集合的均值和离散度转换成伽马分布参数。',
'x̄ 是集合均值，s²ₓ 是成员间样本方差。a、b 控制潜在均值；c、d、e 控制潜在方差。加上式（1）的 δ，共有六个参数。',
'伽马分布满足 μ=κθ、σ²=κθ²，因此 κ=μ²/σ²，θ=σ²/μ。正的均值和方差保证合法的形状与尺度。实际计算设有很小的正值下限，避免数值除零。']),
(5,'3 气候信息如何进入模型',[
'用途：在已有降水订正的基础上，用气候指数调整潜在均值。',
'z 是五维标准化气候指数列向量，依次包含 RMM1、RMM2、Niño 3.4、DMI、SAM；β 是对应的五个系数。βᵀz 是五个“系数乘指数”的和。',
'指数函数始终为正。β 全为零时，乘子为 1，恢复 H0。方差和移动量沿用 H0，因此这是潜在均值的调整，不是把整个降水分布等比例放大。']),
(1,'4 时间累计的一致性',[
'用途：用四个周累计构造所有较长累计。b 是由 w₁、w₂、w₃、w₄ 组成的列向量，T 表示转置。',
'五节点向量依次是四个周和 28 天总量；七节点向量在四周之后依次加入前两周、后两周和 28 天总量。y₅=S₅b，y₇=S₇b。',
'所有降水量均为毫米。这里的 28 天是相同起报日之后的四个完整周，不是自然历月。']),
(2,'5 求和矩阵怎样读',[
'S₇ 有七行四列。前四行组成单位矩阵，只复制四个周值；第五行相加前两周，第六行相加后两周，第七行相加全部四周。',
'例如四周为 1、2、3、4 mm，七节点向量就是 1、2、3、4、3、7、10 mm。S₅ 只保留前四行和最后一行。矩阵元素无量纲。']),
(6,'6 线性协调',[
'用途：把分别拟合、可能不一致的基础预测 ŷ 转换成能够严格相加的预测。',
'S 是 n×4 的求和矩阵，n 为 5 或 7；W 是 n×n 的预测误差协方差，单位 mm²；G 是 4×n 的映射。先用 G 求出四个协调周，再用 S 求出所有累计。−1 表示矩阵求逆。',
'GS=I₄，所以原来已经一致的预测保持不变。在基础预测无偏且协方差给定的条件下，该线性映射具有最小误差方差性质。这一性质不等于样本外 CRPS 一定更低，也不保证降水非负。']),
(7,'7 非负协调',[
'用途：同时满足“降水不小于零”和“各尺度严格相加”。v 是一条已经恢复依赖结构的联合预测样本。',
'argmin 返回使目标函数最小的四周向量 b*。约束 b≥0 对每个周分别成立。目标函数衡量候选一致向量 Sb 与原样本 v 的差异，用 W⁻¹ 加权。最终 v*=Sb*。',
'每个周可以取正值或恰好为零，因此程序检查 16 种可能的边界组合。不能把先做线性协调再逐项截零当成同一个优化问题。此处的非线性投影也不能直接套用式（11）的线性协方差恒等式。']),
(8,'8 对角协方差下的直观解释',[
'用途：解释协调究竟调整了什么。这是五节点、对角 W、没有非负约束时的解析结果。',
'm̂ 是直接预测的 28 天总量，ŵₖ 是第 k 周基础预测；σ²ₘ 和 σ²ₖ 分别是它们的误差方差。s 是四周误差方差之和，只在这条公式中使用这个含义。',
'先算“直接总量减去四周预测之和”，再按各周误差方差分配这项差值。误差越不确定的周允许更大调整。',
'例：四周原来各为 1 mm，直接总量为 8 mm；总量误差方差为 4 mm²，各周为 1 mm²。每周增加 4/8=0.5 mm，协调后四周各为 1.5 mm，总量为 6 mm。这个例子不用于解释完整非对角协方差的权重。']),
(9,'9 CRPS 概率评分',[
'用途：评价一个完整预测分布与实际降水是否相符。y 是观测，x⁽ⁱ⁾ 是第 i 个样本，M 是样本数。',
'第一项是所有样本到观测的平均绝对距离；第二项是样本两两绝对距离平均值的一半。双重求和包含 M² 个有序对，也包含距离为零的自配对。',
'两项一起构成经验分布 CRPS，单位为毫米，越小越好。第二项不能单独理解为“分布越宽越好”。这里只是普通经验 CRPS；fair CRPS 的有限样本修正另有独立随机抽样条件。']),
(10,'10 上尾加权评分',[
'用途：让评分集中反映超过训练期阈值 q 的预测表现。q 是相应季节窗口内的训练降水第 90 百分位数。',
'把样本和观测中低于 q 的值都改成 q，再计算普通 CRPS。这个变换恰好保留阈值以上的积分损失，单位仍为毫米。',
'它不是只挑出观测超过 q 的个例评分。所有起报都参与，因此仍然衡量虚报、漏报以及上尾降水量。']),
(12,'11 从误差协方差预测协调收益',[
'用途：回答一个具体问题：训练误差所显示的依赖关系，预计能让 28 天总量的误差方差比“四周直接相加”减少多少？',
'第一行是核心：预计收益等于 1 减去“协调后的方差 / 周求和的方差”。Vlin,s 是线性协调后的总量误差方差；VBU,s 是四周相加的总量误差方差。两者都是 mm²，比例无量纲。这里 s 是季节标签。',
'第二行计算分子。P=S₇G，是 7×7 的线性协调映射；Ωₛ 是该季节七个基础预测误差的协方差。PΩₛPᵀ 是映射后的误差协方差。[ ]₇₇ 取第七行第七列，因为第七个节点正是 28 天总量。',
'第三行计算分母。Ωₛ,bb 是 Ωₛ 左上角 4×4 的周误差子矩阵，1₄ 是四个元素均为 1 的列向量。这个二次型把四周的方差以及各周之间的协方差全部加起来。它不只是四个对角元素之和。',
'为什么成立：观测本来一致，所以 P 作用在观测上不改变观测。因此协调误差等于 P 乘基础误差，误差协方差相应变成 PΩₛPᵀ。W 用于拟合 P，Ωₛ 用于考察某个季节；它们不必相同。',
'数值例子：若 VBU,s=100 mm²、Vlin,s=80 mm²，则预计收益为 1−80/100=0.20，即预计方差降低 20%。若分子为 120 mm²，则预计收益为 −0.20，表示预计变差 20%。这些是假设数值，用来解释公式，不是论文结果。',
'怎样和实验比较：每个外层验证折都有一个预计比例；图中先对这些比例求平均。实测收益则用各年平均后的 MSE 作比。MSE 还包含偏差的平方，因此与方差并非始终相同。',
'适用边界：公式预测的是固定线性映射下的方差变化，不直接预测 CRPS，也不精确描述非负投影。本文另行检验非负样本均值的表现。已逐项复算 320 个验证年、区域和季节组合，与存档值的最大差异约为 8×10⁻¹⁵。']),
(14,'S1 总量的幂映射',[
'用途：在已有一致预测上，尝试重新校准总量。T 是 28 天总量，s₀ 是正的参考降水尺度，二者均为毫米。a、b 是无量纲参数，与 EMOS 参数只是在不同小节中使用相同字母。',
'T/s₀ 无量纲，exp(a) 调整幅度，exp(b) 给出正指数。a=b=0 时映射不变；T=0 时定义 g(0)=0。该映射保持非负，但不自动保证预测评分改善。']),
(15,'S2 把映射后的总量分回各周',[
'T 是原四周之和。T>0 时，wₖ/T 是第 k 周原来所占比例，乘以 g(T) 得到新周值。四个比例相加为 1，所以四周新值必然加总为 g(T)。',
'T=0 时，四周全为零，避免除以零。原来的周间比例被保留；这个一致性性质与校准是否有效是两件事。']),
(16,'S3 概率混合与权重',[
'用途：组合三个完整联合预测分布。j=1、2、3 分别对应原始集合、BU-clim 和 H1；t 是起报。Pⱼ,ₜ 是分布，wⱼ,ₜ 是其概率权重。',
'xₜ 是包含截距和预报状态的特征列向量，αⱼ 是系数。指数归一化保证权重为正且总和为 1。一个分量的系数固定为零，使参数可识别。这里的 xₜ 不是式（9）的降水样本。',
'对完整一致轨迹赋予概率权重，仍能保持每条轨迹的一致性。概率混合不同于把三个样本值按权重相加。']),
(17,'S4 混合分布的 CRPS',[
'Aⱼ,ₜ 是第 j 个分量样本到观测 yₜ 的平均绝对距离。Bⱼₖ,ₜ 是从分量 j、k 的经验分布独立取值后的平均两两绝对距离；j=k 时也包括分量内部配对。',
'第一项按混合权重汇总观测距离，第二项按两个权重的乘积汇总分量内和分量间距离。A、B 及最终 CRPS 都是毫米。',
'这条公式计算的是给定有限经验分布的精确混合评分，无须重新随机抽样。只有一个权重为 1 时，它退化为该分量的普通经验 CRPS。'])]
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1']:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.25;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.add_paragraph('降水预测论文公式说明','Title')
doc.add_paragraph('对应修订稿主文公式（1）至（11）及补充公式（S1）至（S4）。本说明使用与 Word 正文、LaTeX 相同的规范公式源，帮助理解符号、计算与适用条件。')
md=['# 降水预测论文公式说明'];latex=[]
for key,title,paras in items:
 doc.add_paragraph(title,'Heading 1');p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 if key in [14,15,16,17]:ns['DISPLAY'][key]='S'+str({14:1,15:2,16:3,17:4}[key])
 p._p.append(ns['word_math'](key));md.extend(['## '+title,'$$\n'+ns['EQUATIONS'][key]+'\n$$',*paras])
 for t in paras:doc.add_paragraph(t)
numbered=[(ns['DISPLAY'][key],ns['EQUATIONS'][key]) for key,_,_ in items]
doc.save(ns['OUT']/'JSHESS_equations_explained_CN.docx')
(ns['OUT']/'公式逐条说明.md').write_text('\n\n'.join(md)+'\n',encoding='utf-8')
dest=ROOT/'revision/review_v5/latex_math_review'
body=['\\documentclass[12pt]{article}','\\usepackage[margin=0.8in]{geometry}','\\usepackage{amsmath,amssymb}','\\begin{document}']
for label,eq in numbered:body.append('\\begin{equation}'+eq+'\\tag{'+str(label)+'}\\end{equation}')
body.append('\\end{document}');(dest/'all_equations.tex').write_text('\n'.join(body),encoding='utf-8')
print('Created Chinese guide with all 15 native equations and matching LaTeX formula sheet.')
